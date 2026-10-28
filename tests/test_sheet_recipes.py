#!/usr/bin/env python3
"""Recipes from the owner's Google Sheet.

The sheet is built in the test as an .xlsx laid out like the real one (one
tab per recipe, plus _Template and How to use) and served from 127.0.0.1,
standing in for Google's export link. Today's recipes must come back out of it exactly; a
mistake must cost only the product it is in; a broken or unshared sheet must
change nothing.

Run: QT_QPA_PLATFORM=offscreen python3 tests/test_sheet_recipes.py
"""

import io
import json
import os
import shutil
import sys
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "tests"))

import openpyxl  # noqa: E402

import sheet_recipes as sr  # noqa: E402
from test_updater import World, git  # noqa: E402

with open(os.path.join(ROOT, "recipes.json")) as _fh:
    BASE = json.load(_fh)


def recipe_tabs(cfg=BASE):
    """One tab per recipe, exactly as written into the owner's sheet:
    {tab name: rows}."""
    pap = set(cfg["papain"]["applies_to"])
    tabs = {}
    for p in cfg["products"]:
        fixed = p.get("batch") == "fixed"
        rows = [["ID", p["id"]],
                ["Type", "fixed batch" if fixed else "per kg meat"],
                ["Meat", p["meat"] or ""],
                ["Papain", "yes" if p["id"] in pap else "no"],
                ["Active", "yes"],
                [],
                ["Ingredient", "Grams"]]
        for n, q in p["ingredients"]:
            rows.append([n, float(q) if fixed else round(q * 10, 4)])
        tabs[p["name"]] = rows
    return tabs


def xlsx(tabs, extra_tabs=True):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in tabs.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    if extra_tabs:
        wb.create_sheet("_Template").append(["ID", ""])
        wb.create_sheet("How to use").append(["Recipes for the station"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def row_of(rows, ingredient):
    return next(i for i, r in enumerate(rows) if r and r[0] == ingredient)


# ------------------------------------------------------------- parsing

class TestReadSheet(unittest.TestCase):
    def test_todays_recipes_round_trip_exactly(self):
        parsed = sr.read_sheet(xlsx(recipe_tabs()))
        self.assertEqual(parsed["problems"], {})
        cfg, kept, dropped = sr.merge(BASE, None, parsed)
        self.assertEqual(cfg, BASE, "today's 9 recipes must come back unchanged")
        self.assertEqual((kept, dropped), ([], []))

    def test_label_case_blank_rows_and_ignored_tabs(self):
        tabs = recipe_tabs()
        rows = tabs["Pepper Jerky"]
        rows[0][0], rows[1][0] = "id", "TYPE"
        rows.insert(8, [])                                   # gap in the list
        cfg, _, _ = sr.merge(BASE, None, sr.read_sheet(xlsx(tabs)))
        self.assertEqual(cfg, BASE, "_Template and How to use are not recipes")

    def test_tab_order_is_product_order(self):
        tabs = recipe_tabs()
        reordered = {k: tabs[k] for k in reversed(list(tabs))}
        parsed = sr.read_sheet(xlsx(reordered))
        self.assertEqual([p["name"] for p in parsed["products"]], list(reversed(list(tabs))))

    def test_blank_grams_costs_only_that_recipe(self):
        tabs = recipe_tabs()
        rows = tabs["Pepper Jerky"]
        rows[row_of(rows, "paprika")][1] = None
        parsed = sr.read_sheet(xlsx(tabs))
        self.assertIn("Pepper Jerky", parsed["problems"])
        self.assertIn("paprika", parsed["problems"]["Pepper Jerky"][0])
        self.assertIn("empty", parsed["problems"]["Pepper Jerky"][0])
        self.assertEqual(len(parsed["products"]), 8)

    def test_word_in_grams(self):
        tabs = recipe_tabs()
        rows = tabs["Mughlai Jerky"]
        rows[row_of(rows, "Curd")][1] = "two hundred"
        p = sr.read_sheet(xlsx(tabs))["problems"]["Mughlai Jerky"][0]
        self.assertIn("not a number", p)

    def test_negative_and_duplicate(self):
        tabs = recipe_tabs()
        tabs["Teriyaki Jerky"][row_of(tabs["Teriyaki Jerky"], "honey")][1] = -3
        tabs["Masala Jerky"].append(["salt", 1])            # SALT already there
        probs = sr.read_sheet(xlsx(tabs))["problems"]
        self.assertIn("negative", " ".join(probs["Teriyaki Jerky"]))
        self.assertIn("twice", " ".join(probs["Masala Jerky"]))

    def test_papain_yes_on_fixed_batch(self):
        tabs = recipe_tabs()
        tabs["Piri Piri Masala"][3][1] = "yes"
        probs = sr.read_sheet(xlsx(tabs))["problems"]
        self.assertIn("fixed batch", probs["Piri Piri Masala"][0])

    def test_bad_or_empty_type(self):
        tabs = recipe_tabs()
        tabs["Vinegar Bath"][1][1] = "kg"
        tabs["Mughlai Jerky"][1][1] = ""
        probs = sr.read_sheet(xlsx(tabs))["problems"]
        self.assertIn("Vinegar Bath", probs)
        self.assertIn("Type is empty", probs["Mughlai Jerky"])

    def test_missing_header_row(self):
        tabs = recipe_tabs()
        tabs["Vinegar Bath"][6] = ["Item", "Weight"]
        probs = sr.read_sheet(xlsx(tabs))["problems"]
        self.assertIn("header", probs["Vinegar Bath"][0])

    def test_inactive_recipe_left_out(self):
        tabs = recipe_tabs()
        tabs["Mughlai Jerky"][4][1] = "no"
        tabs["Mughlai Jerky"][row_of(tabs["Mughlai Jerky"], "Curd")][1] = "broken"
        parsed = sr.read_sheet(xlsx(tabs))
        self.assertEqual(parsed["problems"], {})
        self.assertNotIn("mughlai_jerky", [p["id"] for p in parsed["products"]])
        self.assertEqual(parsed["inactive"], ["Mughlai Jerky"])
        cfg, _, _ = sr.merge(BASE, None, parsed)
        self.assertNotIn("mughlai_jerky", cfg["papain"]["applies_to"])

    def test_new_recipe_tab_gets_id_from_name(self):
        tabs = recipe_tabs()
        tabs["Smoky BBQ Jerky"] = [["ID", ""], ["Type", "per kg meat"], ["Meat", "Buffalo"],
                                   ["Papain", "yes"], ["Active", "yes"], [],
                                   ["Ingredient", "Grams"], ["bbq sauce", 120], ["salt", 8]]
        cfg, _, _ = sr.merge(BASE, None, sr.read_sheet(xlsx(tabs)))
        new = cfg["products"][-1]
        self.assertEqual(new["id"], "smoky_bbq_jerky")
        self.assertEqual(new["meat"], "Buffalo")
        self.assertEqual(new["ingredients"], [["bbq sauce", 12.0], ["salt", 0.8]])
        self.assertIn("smoky_bbq_jerky", cfg["papain"]["applies_to"])

    def test_renamed_tab_keeps_its_id(self):
        tabs = recipe_tabs()
        tabs = {("Black Pepper Jerky" if k == "Pepper Jerky" else k): v for k, v in tabs.items()}
        cfg, _, _ = sr.merge(BASE, None, sr.read_sheet(xlsx(tabs)))
        p = next(p for p in cfg["products"] if p["id"] == "pepper_jerky")
        self.assertEqual(p["name"], "Black Pepper Jerky")

    def test_copied_tab_with_old_id_is_a_mistake_in_the_copy_only(self):
        tabs = recipe_tabs()
        copy = json.loads(json.dumps(tabs["Teriyaki Jerky"]))
        tabs["Copy of Teriyaki Jerky"] = copy               # ID still teriyaki_jerky
        parsed = sr.read_sheet(xlsx(tabs))
        self.assertIn("Copy of Teriyaki Jerky", parsed["problems"])
        self.assertIn("empty the ID", parsed["problems"]["Copy of Teriyaki Jerky"][0])
        cfg, kept, dropped = sr.merge(BASE, BASE, parsed)
        self.assertEqual([p["id"] for p in cfg["products"]].count("teriyaki_jerky"), 1)
        self.assertEqual(dropped, ["Copy of Teriyaki Jerky"])

    def test_no_recipe_tabs(self):
        with self.assertRaises(sr.SheetError) as e:
            sr.read_sheet(xlsx({}))
        self.assertIn("no recipe tabs", str(e.exception))

    def test_signin_page_means_not_shared(self):
        with self.assertRaises(sr.SheetError) as e:
            sr.read_sheet(b"<!DOCTYPE html><html>Sign in</html>")
        self.assertIn("Anyone with the link", str(e.exception))

    def test_merge_keeps_previous_for_problem_recipe(self):
        prev = json.loads(json.dumps(BASE))
        tabs = recipe_tabs()
        tabs["Pepper Jerky"][row_of(tabs["Pepper Jerky"], "paprika")][1] = None
        tabs["Half Done Jerky"] = [["ID", ""], ["Type", "per kg meat"], [], ["Ingredient", "Grams"]]
        parsed = sr.read_sheet(xlsx(tabs))
        cfg, kept, dropped = sr.merge(BASE, prev, parsed)
        self.assertEqual(kept, ["Pepper Jerky"])
        self.assertEqual(dropped, ["Half Done Jerky"])
        pepper = next(p for p in cfg["products"] if p["id"] == "pepper_jerky")
        self.assertEqual(pepper, next(p for p in prev["products"] if p["id"] == "pepper_jerky"))
        self.assertIn("pepper_jerky", cfg["papain"]["applies_to"])

    def test_export_url(self):
        url = ("https://docs.google.com/spreadsheets/d/"
               "1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789-_xyz/edit?gid=0#gid=0")
        self.assertEqual(sr.export_url(url),
                         "https://docs.google.com/spreadsheets/d/"
                         "1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789-_xyz/export?format=xlsx")
        self.assertEqual(sr.export_url("http://127.0.0.1:1/s.xlsx"), "http://127.0.0.1:1/s.xlsx")


# ------------------------------------------------------- the updater, live

class FakeGoogle:
    """Serves the sheet like Google's export link: the .xlsx, a sign-in
    page (not shared), or 403."""

    def __init__(self):
        self.body, self.status = xlsx(recipe_tabs()), 200
        me = self

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def do_GET(self):
                self.send_response(me.status)
                self.send_header("Content-Length", str(len(me.body)))
                self.end_headers()
                self.wfile.write(me.body)

        self.server = ThreadingHTTPServer(("127.0.0.1", 0), H)
        self.url = f"http://127.0.0.1:{self.server.server_address[1]}/export.xlsx"
        threading.Thread(target=self.server.serve_forever, daemon=True).start()

    def stop(self):
        self.server.shutdown()
        self.server.server_close()


class TestUpdaterFromSheet(unittest.TestCase):
    def setUp(self):
        self.w = World()
        self.g = FakeGoogle()
        self.conf = os.path.join(self.w.home, ".config", "mincecraft")
        os.makedirs(self.conf)
        self.set_link(self.g.url)

    def tearDown(self):
        self.g.stop()
        self.w.close()

    def set_link(self, url):
        with open(os.path.join(self.conf, "recipes.env"), "w") as fh:
            fh.write(f"RECIPES_SHEET_URL={url}\n")

    def run_u(self, *args):
        return self.w.updater("--config-dir", self.conf, *args)

    def test_first_sync_equals_repo_recipes(self):
        rc, out = self.run_u("--recipes")
        self.assertEqual(rc, 0, out)
        local = self.w.local()
        sync = local.pop("_sync")
        self.assertEqual(local, BASE)
        self.assertEqual(sync["source"], "google-sheet")
        self.assertTrue(sync["commit"].startswith("sheet-"))
        st = self.w.status()["recipes"]
        self.assertTrue(st["ok"])
        self.assertEqual(st["source"], "google-sheet")
        self.assertTrue(self.w.clean())

    def test_owner_changes_grams_and_pi_follows_once(self):
        self.run_u("--recipes")
        tabs = recipe_tabs()
        tabs["Teriyaki Jerky"][row_of(tabs["Teriyaki Jerky"], "soy sauce")][1] = 50
        self.g.body = xlsx(tabs)
        self.run_u("--recipes")
        p = next(p for p in self.w.local()["products"] if p["id"] == "teriyaki_jerky")
        self.assertEqual(dict(p["ingredients"])["soy sauce"], 5.0)
        path = os.path.join(self.w.data, "recipes.json")
        m = os.stat(path).st_mtime_ns
        self.run_u("--recipes")
        self.assertEqual(m, os.stat(path).st_mtime_ns, "unchanged sheet: no rewrite")

    def test_mistake_keeps_previous_recipe(self):
        self.run_u("--recipes")
        tabs = recipe_tabs()
        tabs["Pepper Jerky"][row_of(tabs["Pepper Jerky"], "paprika")][1] = None
        tabs["Teriyaki Jerky"][row_of(tabs["Teriyaki Jerky"], "honey")][1] = 44
        self.g.body = xlsx(tabs)
        self.run_u("--recipes")
        local = self.w.local()
        ter = next(p for p in local["products"] if p["id"] == "teriyaki_jerky")
        pep = next(p for p in local["products"] if p["id"] == "pepper_jerky")
        self.assertEqual(dict(ter["ingredients"])["honey"], 4.4)
        self.assertEqual(pep, next(p for p in BASE["products"] if p["id"] == "pepper_jerky"))
        st = self.w.status()["recipes"]
        self.assertEqual([x["sheet"] for x in st["skipped_sheets"]], ["Pepper Jerky"])
        self.assertTrue(st["skipped_sheets"][0]["kept_previous"])

    def test_broken_sheet_rejected_once(self):
        self.run_u("--recipes")
        good = self.w.local()
        self.g.body = xlsx({})                     # only _Template and How to use
        rc, out = self.run_u("--recipes")
        self.assertEqual(rc, 0)
        self.assertEqual(self.w.local(), good)
        st = self.w.status()["recipes"]
        self.assertIs(st["ok"], False)
        self.assertIn("no recipe tabs", st["rejected"]["why"])
        self.assertIn("REJECTED", out)
        rc, out = self.run_u("--recipes")
        self.assertNotIn("REJECTED", out)

    def test_not_shared(self):
        self.run_u("--recipes")
        good = self.w.local()
        self.g.body = b"<html>Sign in to continue to Sheets</html>"
        self.run_u("--recipes")
        self.assertEqual(self.w.local(), good)
        self.assertIn("Anyone with the link", self.w.status()["recipes"]["rejected"]["why"])
        self.g.status, self.g.body = 403, b"forbidden"
        self.run_u("--recipes")
        self.assertEqual(self.w.local(), good)

    def test_offline(self):
        self.run_u("--recipes")
        good = self.w.local()
        self.g.stop()
        self.g = FakeGoogle()                    # a new port: the old one is closed
        self.set_link("http://127.0.0.1:9/export.xlsx")
        rc, _ = self.run_u("--recipes")
        self.assertEqual(rc, 0)
        self.assertEqual(self.w.local(), good)
        self.assertIs(self.w.status()["recipes"]["online"], False)

    def test_boot_syncs_sheet_even_if_github_is_down(self):
        git(self.w.pi, "remote", "set-url", "origin", "/nonexistent/repo.git")
        rc, _ = self.run_u("--boot")
        self.assertEqual(rc, 0)
        self.assertEqual(self.w.local()["_sync"]["source"], "google-sheet")
        self.assertIn("GitHub", self.w.status()["code"]["last_error"])

    def test_no_link_falls_back_to_workbook(self):
        self.set_link("PASTE-THE-GOOGLE-SHEET-LINK-HERE")
        self.run_u("--recipes")
        self.assertEqual(self.w.local()["_sync"]["source"], "origin/main")


# ----------------------------------------------------------- the station

class TestStationWording(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PyQt5.QtWidgets import QApplication
        cls.app = QApplication.instance() or QApplication(sys.argv[:1])

    def test_home_names_the_sheet_and_records_source(self):
        from scale import BatchLog, Config, DailyRatio, ScaleState
        from panel import Panel, Step
        d = tempfile.mkdtemp(prefix="_test_sheet_panel_")
        try:
            cfg = json.loads(json.dumps(BASE))
            cfg["_sync"] = {"source": "google-sheet", "commit": "sheet-abc123def456",
                            "applied_at": "2026-10-27T16:00:00+0530"}
            local = os.path.join(d, "recipes.json")
            json.dump(cfg, open(local, "w"))
            status = os.path.join(d, "update.json")
            json.dump({"recipes": {"source": "google-sheet", "ok": True,
                                   "commit": "sheet-abc123def456",
                                   "skipped_sheets": [{"sheet": "Pepper Jerky"}]}},
                      open(status, "w"))
            c = Config.load(local)
            log = os.path.join(d, "b.jsonl")
            win = Panel(ScaleState(), c, BatchLog(log),
                        DailyRatio(os.path.join(d, "dd.json"), c),
                        recipes_path=local, status_path=status)
            win.refresh_consumption = lambda: None
            win.check_recipes()
            tone, text = win.update_notice()
            self.assertEqual(tone, "warn")
            self.assertIn("recipe sheet", text)
            self.assertIn("Pepper Jerky", text)
            json.dump({"recipes": {"source": "google-sheet", "ok": True,
                                   "commit": "sheet-abc123def456", "online": False}},
                      open(status, "w"))
            win.check_recipes()
            self.assertIn("not reachable", win.update_notice()[1])
            json.dump({"recipes": {"source": "google-sheet", "ok": True,
                                   "commit": "sheet-abc123def456"}}, open(status, "w"))
            win.check_recipes()
            self.assertIn("Google Sheet", win.update_notice()[1])
            win.start_batch()
            win.st.product = "piri_piri_masala"
            win.st.steps = [Step(n, None, t) for n, t in c.targets_for("piri_piri_masala")]
            for s in win.st.steps:
                s.actual = s.target
            win.record_batch()
            rec = json.loads(open(log).read().splitlines()[-1])
            self.assertEqual(rec["recipes_source"], "google-sheet")
            self.assertEqual(rec["recipes_commit"], "sheet-abc123def456")
            win.close()
        finally:
            shutil.rmtree(d, ignore_errors=True)


if __name__ == "__main__":
    unittest.main(verbosity=1)
