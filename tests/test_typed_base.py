#!/usr/bin/env python3
"""Typing the meat weight instead of weighing it.

From the capture screen a supervisor can press TYPE WEIGHT, enter the PIN and
type the meat's weight in kilograms. The meat never goes on the scale, so the
container swap is skipped and the batch goes straight to recipe review. The
record says the weight was typed.

Run: QT_QPA_PLATFORM=offscreen python3 tests/test_typed_base.py
"""

import json
import os
import sys
import tempfile
import time

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import openpyxl  # noqa: E402
from PyQt5.QtCore import QTimer  # noqa: E402
from PyQt5.QtWidgets import QApplication, QPushButton  # noqa: E402

import consumption  # noqa: E402
from scale import (MAIN, SMALL, BatchLog, Config, DailyRatio,  # noqa: E402
                   ScaleState, SimScale, start_reader)
from panel import Panel, Step, scale_for  # noqa: E402

SHOTS = os.path.join(ROOT, "screenshots")
LOG = os.path.join(ROOT, "_test_typed.jsonl")
DAILY = os.path.join(ROOT, "_test_typed_daily.json")
PID = "pepper_jerky"
CONTAINER_G = 640.0

checks = []


def check(label, cond):
    checks.append((label, bool(cond)))
    print(("  PASS  " if cond else "  FAIL  ") + label)


def pump(app, seconds=0.3):
    end = time.time() + seconds
    while time.time() < end:
        app.processEvents()
        time.sleep(0.01)


def settle(app, state, seconds=8):
    end = time.time() + seconds
    run = 0
    while time.time() < end:
        run = run + 1 if state.snapshot()["stable"] else 0
        if run >= 3:
            pump(app, 0.4)
            return True
        pump(app, 0.1)
    return False


def last_record():
    with open(LOG, encoding="utf-8") as fh:
        return json.loads(fh.read().strip().splitlines()[-1])


def type_keys(pad, text):
    pad.clear()
    for k in text:
        pad.press(k)


def pin_then(app, keys, cancel=False):
    """Press keys on the PIN dialog once it is up (it runs modally)."""
    def go():
        dlg = app.activeModalWidget()
        if dlg is None:
            QTimer.singleShot(50, go)
            return
        for k in keys:
            dlg.key(k)
        if cancel and dlg.isVisible():
            dlg.reject()
    QTimer.singleShot(100, go)


def cleanup():
    for f in (LOG, DAILY):
        if os.path.exists(f):
            os.remove(f)


def main():
    os.makedirs(SHOTS, exist_ok=True)
    cleanup()
    cfg = Config.load()
    app = QApplication.instance() or QApplication(sys.argv[:1])

    # --------------------------------------------------------- the limits
    print("--- the limits on a typed weight ---")
    floor = cfg.min_base_for(PID)
    check("nothing typed is refused", cfg.typed_base_problem(PID, None))
    check("zero is refused", cfg.typed_base_problem(PID, 0))
    check(f"below the recipe minimum ({floor:.0f} g) is refused",
          cfg.typed_base_problem(PID, floor - 1))
    check("the minimum itself is fine", cfg.typed_base_problem(PID, floor) is None)
    check("a slipped decimal point (325 kg) is refused",
          "decimal" in (cfg.typed_base_problem(PID, 325000) or ""))
    check("3.25 kg is fine", cfg.typed_base_problem(PID, 3250) is None)
    tery = cfg.min_base_for("teriyaki_jerky")
    check(f"Teriyaki keeps its own higher minimum ({tery:.0f} g)",
          cfg.typed_base_problem("teriyaki_jerky", tery - 1)
          and cfg.typed_base_problem("teriyaki_jerky", tery) is None)

    # -------------------------------------------- with no scale at all
    print("--- available when the scale is dead ---")
    dead = Panel(ScaleState(), cfg, BatchLog(LOG), DailyRatio(DAILY, cfg))
    dead.resize(1024, 600)
    dead.show()
    dead.choose_product(PID)
    pump(app, 0.5)
    cap = dead.screens["CAPTURE"]
    check("CAPTURE is off with no scale", not cap.cap_btn.isEnabled())
    check("TYPE WEIGHT is still there and on",
          cap.type_btn.isVisible() and cap.type_btn.isEnabled())
    check("the hint points to it", "TYPE WEIGHT" in cap.hint.text())
    dead.close()

    # ----------------------------------------- the real flow at 800x480
    print("--- typed weight, end to end, on the 7-inch screen ---")
    state = ScaleState(stable_band_g=2 * cfg.main.division_g)
    sim = SimScale(division_g=cfg.division_g)
    _t, stop = start_reader(state, sim=sim)
    win = Panel(state, cfg, BatchLog(LOG), DailyRatio(DAILY, cfg), sim=sim,
                scale=scale_for(800, 480))
    win.resize(800, 480)
    win.show()
    seen = []
    orig_show = win.show_screen

    def spy(key):
        seen.append(key)
        orig_show(key)
    win.show_screen = spy
    pump(app, 1.0)

    win.screens["HOME"].start_btn.click()
    pump(app, 0.3)
    win.choose_product(PID)
    pump(app, 0.4)
    check("a meat recipe opens on the capture screen", win.current == "CAPTURE")
    cap = win.screens["CAPTURE"]
    check("TYPE WEIGHT fits beside the other buttons at 800x480",
          cap.type_btn.isVisible()
          and cap.type_btn.width() >= cap.type_btn.sizeHint().width() - 1
          and cap.type_btn.geometry().right() < win.width())

    pin_then(app, list("9999") + ["OK"], cancel=True)
    cap.type_btn.click()
    pump(app, 0.4)
    check("a wrong PIN leaves you on the capture screen", win.current == "CAPTURE")

    pin_then(app, list(cfg.pin) + ["OK"])
    cap.type_btn.click()
    pump(app, 0.4)
    check("the right PIN opens the typing screen", win.current == "TYPED")
    typed = win.screens["TYPED"]
    check("it asks in kg and says the weight will be recorded as typed",
          "kg" in typed.headline.text() and "typed" in typed.limits.text())
    check("nothing typed yet — USE is off", not typed.use_btn.isEnabled())

    type_keys(typed.pad, "325")
    check("325 (kg) is refused with a decimal-point hint",
          not typed.use_btn.isEnabled() and "decimal" in typed.guide.text())
    type_keys(typed.pad, "0.3")
    check("0.3 kg is refused as below the minimum",
          not typed.use_btn.isEnabled() and "below" in typed.guide.text())
    type_keys(typed.pad, "3.25")
    check("3.25 kg is accepted, shown as 3,250 g",
          typed.use_btn.isEnabled() and "3,250 g" in typed.guide.text())
    check("and it says to put the empty container on",
          "container" in typed.guide.text())
    win.grab().save(os.path.join(SHOTS, "32-typed-weight.png"))

    # The empty ingredient container goes on, as the screen says.
    sim.set(CONTAINER_G)
    settle(app, state)
    typed.use_btn.click()
    pump(app, 0.5)
    check("USE goes straight to recipe review", win.current == "REVIEW")
    check("never through the container swap", "SWAP" not in seen)
    check("the meat weight is the typed one", abs(win.st.base_wt - 3250) < 1e-9)
    want = {n: t for n, t in cfg.targets_for(PID, 3250)}
    check("every target computed from 3.25 kg",
          {s.name: s.target for s in win.st.steps} == want)
    rev = win.screens["REVIEW"]
    check("review says the meat was typed", "typed" in rev.crumb.text())
    check("START reminds the container must be on",
          "CONTAINER ON" in rev.start_btn.text())
    check("the status bar says (typed)", "(typed)" in win.info_lbl.text())
    check("papain preview uses the typed weight (buffalo 19.5 g)",
          any("if buffalo" in (rev.table.item(r, 1).text() if rev.table.item(r, 1) else "")
              and "19.5" in rev.table.item(r, 3).text()
              for r in range(rev.table.rowCount())))
    win.grab().save(os.path.join(SHOTS, "33-review-typed.png"))

    rev.start_btn.click()
    pump(app, 0.5)
    check("the ingredients start from the container, not from zero",
          abs((win.st.start_g or 0) - CONTAINER_G) <= 2)

    for i, s in enumerate(list(win.st.steps)):
        if s.scale == SMALL:
            if win.current != "MANUAL":
                break
            sim.add(s.target)
            settle(app, state)
            win.screens["MANUAL"].confirm_btn.click()
            pump(app, 0.5)
            continue
        if win.current != "ADD":
            break
        sim.add(s.target)
        end = time.time() + 15
        while time.time() < end and win.current == "ADD" and win.st.idx == i:
            pump(app, 0.1)
    check("every ingredient in, then the meat question", win.current == "MEAT")
    win.screens["MEAT"].buttons["buffalo"].click()
    pump(app, 0.4)
    pap = win.st.steps[-1]
    check("buffalo papain from the typed weight: 6 g/kg x 3.25 = 19.5 g",
          pap.name == cfg.papain_name and abs(pap.target - 19.5) < 1e-9)
    sim.add(pap.target)
    end = time.time() + 15
    while time.time() < end and win.current == "ADD":
        pump(app, 0.1)
    check("batch complete", win.current == "DONE")

    rec = last_record()
    check("record: meat weight 3250 g, source 'typed'",
          rec["base_weight_g"] == 3250 and rec["base_weight_source"] == "typed")
    check("record: per-meat batch, buffalo",
          rec["batch"] == "per_meat" and rec["meat_type"] == "buffalo")
    recon = rec["reconciliation"]
    check("it reconciles from the container "
          f"(expected {recon.get('expected_g')}, saw {recon.get('observed_g')})",
          recon.get("available") and recon.get("ok"))

    # ------------------------------ weighing again afterwards is unaffected
    print("--- the next batch, weighed normally ---")
    win.new_batch()
    pump(app, 0.3)
    check("a new batch forgets the typed flag", win.st.base_typed is False)
    sim.zero()
    win.screens["HOME"].start_btn.click()
    pump(app, 0.3)
    win.choose_product(PID)
    sim.set(3000)
    settle(app, state)
    win.screens["CAPTURE"].cap_btn.click()
    pump(app, 0.3)
    check("CAPTURE still goes to the container swap", win.current == "SWAP")
    check("and the batch is not marked typed", win.st.base_typed is False)
    check("review START is the plain one",
          (win.pick_product(PID) or True)
          and win.screens["REVIEW"].start_btn.text() == "START ADDING ▶")
    for s in win.st.steps:
        s.actual, s.verified = s.target, True
    win.st.idx = len(win.st.steps) - 1
    win.next_step()
    pump(app, 0.3)
    win.screens["MEAT"].buttons["other"].click()
    pump(app, 0.3)
    check("a weighed batch records source 'scale'",
          last_record()["base_weight_source"] == "scale")

    # ---------------------------------------- fixed batch: no meat, no source
    win.new_batch()
    win.st.product = "piri_piri_masala"
    win.st.steps = [Step(n, None, t, scale=cfg.scale_for(t) or MAIN)
                    for n, t in cfg.targets_for("piri_piri_masala")]
    for s in win.st.steps:
        s.actual, s.verified = s.target, True
    win.st.idx = len(win.st.steps) - 1
    win.next_step()
    pump(app, 0.3)
    check("a fixed batch records no meat-weight source",
          last_record()["base_weight_source"] is None)

    # ------------------------------------------- the consumption workbook
    out = os.path.join(tempfile.mkdtemp(), "consumption.xlsx")
    consumption.build(LOG, out)
    ws = openpyxl.load_workbook(out)["Batches"]
    check("Batches sheet has a 'Meat weight' column",
          ws.cell(row=4, column=13).value == "Meat weight")
    col = [ws.cell(row=r, column=13).value for r in range(5, 8)]
    check(f"typed / scale / blank for the three batches (got {col})",
          col == ["typed", "scale", ""] or col == ["typed", "scale", None])

    stop.set()
    win.close()
    cleanup()
    failed = [c for c, ok in checks if not ok]
    print(f"\n{len(checks) - len(failed)}/{len(checks)} checks passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
